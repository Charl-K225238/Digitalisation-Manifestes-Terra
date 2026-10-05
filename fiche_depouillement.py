"""
Fiche de dépouillement (intégrée depuis la mini-app « Manifeste parser v5 »).

À partir des B/L déjà extraits par `manifest_parser.parse_manifest`, compte les
unités par catégorie de la fiche : USED / NEW / 20' / 40' / >50m3 / BOLSTER /
DIVERS. La classification suit le libellé de catégorie Grimaldi (pas le volume) :
  - Used/New Small Van, Big Van, Car → USED / NEW ;
  - Used/New LM RoRo → USED / NEW (grand van) ;
  - LM RoRo, LM Cargo, tracteurs, remorques, engins, « bébés au dos » → >50m3 ;
  - Conteneurs → 20' / 40' ; Bolster → BOLSTER ; le reste → DIVERS.

Pas de dépendance Streamlit : testable seul.
"""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from manifest_parser import MODEL_YEAR_RE, extract_marque_modele, item_status

CATS = ["USED", "NEW", "20'", "40'", ">50m3", "BOLSTER", "DIVERS"]

SUBTYPE_LABELS = {
    "USED_SMALL_VAN": "SUV / Petit Van USED",
    "USED_BIG_VAN": "Grand Van / Pickup USED",
    "USED_CAR": "Voiture USED",
    "NEW_SMALL_VAN": "SUV / Petit Van NEW",
    "NEW_BIG_VAN": "Grand Van / Pickup NEW",
    "NEW_CAR": "Voiture NEW",
    "CONT_20": "Container 20'",
    "CONT_40HC": "Container 40' HC",
    "CONT_40OT": "Container 40' OT",
    "MAFI": "LM RoRo / >50m3",
    "BOLSTER": "BOLSTER",
    "DIVERS": "DIVERS",
}
SUBTYPE_CAT = {
    "USED_SMALL_VAN": "USED", "USED_BIG_VAN": "USED", "USED_CAR": "USED",
    "NEW_SMALL_VAN": "NEW", "NEW_BIG_VAN": "NEW", "NEW_CAR": "NEW",
    "CONT_20": "20'", "CONT_40HC": "40'", "CONT_40OT": "40'",
    "MAFI": ">50m3", "BOLSTER": "BOLSTER", "DIVERS": "DIVERS",
}

# Ordre significatif : le premier motif qui correspond gagne.
_PATTERNS = [
    (r'used\s+small\s+van', 'USED_SMALL_VAN'),
    (r'used\s+big\s+van', 'USED_BIG_VAN'),
    (r'used\s+car', 'USED_CAR'),
    (r'new\s+small\s+van', 'NEW_SMALL_VAN'),
    (r'new\s+big\s+van', 'NEW_BIG_VAN'),
    (r'new\s+car', 'NEW_CAR'),
    (r'used\s+lm\s+roro', 'USED_BIG_VAN'),
    (r'new\s+lm\s+roro', 'NEW_BIG_VAN'),
    (r'lm\s+roro', 'MAFI'),
    (r'lm\s+cargo', 'MAFI'),
    (r'bolster', 'BOLSTER'),
    (r'20\s*ft', 'CONT_20'),
    (r'40\s*ft\.?\s*(open\s*top|flat\s*rack)', 'CONT_40OT'),
    (r'40\s*ft', 'CONT_40HC'),
    (r'tractor|trailer|construction\s+equip|heavy\s+equip|excavat|bulldozer|'
     r'forklift|crane|generator|compressor|truck|\bbus\b|ambulance', 'MAFI'),
    (r'boat|yacht|motor\s*cycle', 'DIVERS'),
]
_PATTERNS = [(re.compile(p, re.I), c) for p, c in _PATTERNS]
_BARE_VEHICLE = [  # véhicule sans « used/new » dans le libellé : statut à chercher
    (re.compile(r'small\s+van', re.I), "SMALL_VAN"),
    (re.compile(r'big\s+van', re.I), "BIG_VAN"),
    (re.compile(r'\bcar(?:\(s\))?\b', re.I), "CAR"),
]


def classify_item(type_raw, is_mafi=False, piggyback=False, context=""):
    """Retourne (sous_type, fiabilité). Fiabilité : HIGH (libellé explicite),
    MEDIUM (statut Used/New déduit du contexte B/L), LOW (à vérifier)."""
    if piggyback or is_mafi:
        return "MAFI", "HIGH"
    t = type_raw or ""
    for rx, sub in _PATTERNS:
        if rx.search(t):
            return sub, "HIGH"
    for rx, kind in _BARE_VEHICLE:
        if rx.search(t):
            st = item_status(t, context)
            if st == "Neuf":
                return f"NEW_{kind}", "MEDIUM"
            if st == "Usager":
                return f"USED_{kind}", "MEDIUM"
            return f"USED_{kind}", "LOW"
    return "DIVERS", ("LOW" if t.strip() else "MEDIUM")


def _split_vessel(vv):
    m = re.match(r'^(.*?)\s*:\s*(\S+)$', vv or "")
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return (vv or "").strip() or "(navire non détecté)", ""


def _per_unit(total, qty, nd):
    try:
        return round(float(total) / qty, nd) if total else None
    except (TypeError, ValueError, ZeroDivisionError):
        return None


def build_units(records):
    """Une ligne par unité physique (qty éclatée). Liste de dicts."""
    units = []
    for r in records:
        navire, voyage = _split_vessel(r.get("vessel_voyage"))
        ctx = " | ".join(dict.fromkeys(r.get("raw_desc_lines", [])))
        for it in r.get("items", []):
            raw = it.get("type_raw", "")
            sub, conf = classify_item(raw, it.get("is_mafi", False),
                                      it.get("_piggyback", False), ctx)
            qty = max(int(it.get("qty") or 1), 1)
            marque, modele = extract_marque_modele(raw)
            if not marque:
                marque, modele = extract_marque_modele(ctx)
            ym = MODEL_YEAR_RE.search(raw + " " + ctx)
            annee = (ym.group(1) or ym.group(2)) if ym else ""
            chassis = it.get("chassis") or []
            conts = it.get("container_no") or []
            w, c = _per_unit(it.get("weight"), qty, 1), _per_unit(it.get("cbm"), qty, 3)
            for i in range(qty):
                units.append({
                    "Navire": navire, "Voyage": voyage,
                    "BL": r.get("bl_number", ""),
                    "Catégorie": SUBTYPE_CAT[sub], "Sous-type": SUBTYPE_LABELS[sub],
                    "Identification manifeste": raw,
                    "Marque/Modèle": " ".join(x for x in (marque, modele) if x),
                    "N° conteneur": conts[i] if i < len(conts) else "",
                    "Châssis/VIN": chassis[i] if i < len(chassis) else "",
                    "Année": annee, "Poids (kg)": w, "CBM": c,
                    "Fiabilité": conf,
                })
    return units


def count_by_category(units):
    counts = {c: 0 for c in CATS}
    for u in units:
        counts[u["Catégorie"]] += 1
    return counts


def count_by_subtype(units):
    out = {}
    for u in units:
        k = (u["Catégorie"], u["Sous-type"])
        out[k] = out.get(k, 0) + 1
    return out


# ── Export Excel (design repris de la mini-app v5) ─────────────────────────
_NAVY, _ACCENT = "1B2B4B", "2563EB"
_ROW_COLORS = {
    "USED": ("FEF3C7", "92400E"), "NEW": ("D1FAE5", "065F46"),
    "20'": ("E0E7FF", "3730A3"), "40'": ("F3E8FF", "6B21A8"),
    ">50m3": ("FCE7F3", "9D174D"), "BOLSTER": ("FEE2E2", "991B1B"),
    "DIVERS": ("F3F4F6", "6B7280"),
}
_FONT = "Segoe UI"
_thin = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def build_fiche_excel(units, navire="", voyage=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fiche"
    counts = count_by_category(units)
    total = sum(counts.values())

    ws.merge_cells("A1:H1")
    ws["A1"] = f"FICHE DE DÉPOUILLEMENT — {navire} {voyage}".strip()
    ws["A1"].font = Font(name=_FONT, size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = _fill(_NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    heads = CATS + ["TOTAL"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=3, column=j, value=h)
        bg, fg = _ROW_COLORS.get(h, (_NAVY, "FFFFFF"))
        c.fill = _fill(bg if h != "TOTAL" else _NAVY)
        c.font = Font(name=_FONT, bold=True, color=fg if h != "TOTAL" else "FFFFFF")
        c.alignment = Alignment(horizontal="center")
        c.border = _BORDER
        v = ws.cell(row=4, column=j, value=counts.get(h, total))
        v.font = Font(name=_FONT, size=16, bold=True)
        v.alignment = Alignment(horizontal="center")
        v.border = _BORDER

    ws["A6"] = "Détail par sous-type"
    ws["A6"].font = Font(name=_FONT, bold=True, color=_ACCENT)
    r = 7
    for j, h in enumerate(["Catégorie", "Sous-type", "Nombre"], 1):
        c = ws.cell(row=r, column=j, value=h)
        c.font = Font(name=_FONT, bold=True, color="FFFFFF")
        c.fill = _fill(_NAVY)
        c.border = _BORDER
    for (cat, sub), n in sorted(count_by_subtype(units).items(),
                                key=lambda kv: (CATS.index(kv[0][0]), kv[0][1])):
        r += 1
        bg, fg = _ROW_COLORS[cat]
        for j, val in enumerate([cat, sub, n], 1):
            c = ws.cell(row=r, column=j, value=val)
            c.fill, c.border = _fill(bg), _BORDER
            c.font = Font(name=_FONT, color=fg)
    for j in range(1, 9):
        ws.column_dimensions[get_column_letter(j)].width = 22 if j == 2 else 13

    wd = wb.create_sheet("Détail")
    cols = ["BL", "Catégorie", "Sous-type", "Identification manifeste", "Marque/Modèle",
            "N° conteneur", "Châssis/VIN", "Année", "Poids (kg)", "CBM", "Fiabilité"]
    for j, h in enumerate(cols, 1):
        c = wd.cell(row=1, column=j, value=h if h != "BL" else "N° B/L")
        c.font = Font(name=_FONT, bold=True, color="FFFFFF")
        c.fill = _fill(_NAVY)
        c.border = _BORDER
    for i, u in enumerate(units, 2):
        bg, fg = _ROW_COLORS[u["Catégorie"]]
        for j, h in enumerate(cols, 1):
            c = wd.cell(row=i, column=j, value=u[h])
            c.fill, c.border = _fill(bg), _BORDER
            c.font = Font(name=_FONT, color=fg)
    widths = [18, 11, 24, 42, 22, 16, 22, 8, 11, 9, 11]
    for j, w in enumerate(widths, 1):
        wd.column_dimensions[get_column_letter(j)].width = w
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(units), 1) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
