"""Calcul du bloc mensuel « Reporting RORO & TEUS », contrôles de cohérence,
lecture du référentiel (rapport existant) et export Excel.

Règles validées avec le service (06/10/2026) :
- Escales      = navires listés dans le classeur volumes du mois.
- TEU          = 20' + 2 × 40', pleins + vides, import + export. MAFI et
                 bolsters exclus (= « TOTAL EVP » du classeur).
- RORO         = TOTAL VEHICULE, import + export.
- Neufs/usagés = DT VEHICULES NEUFS / USAGES, import + export.
- Tranches     = extrait PAA (rubriques UNL/LOAD, shifting et bord à bord
                 exclus). Lo/Lo = navires de type COMBONG dans le PAA.
- Hinterland   = véhicules en transit ; le total vient du classeur, la
                 répartition par tranche est saisie (pas dans les sources).
Toute valeur peut être corrigée à la main ; la valeur calculée est gardée.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

import pandas as pd

from stats_flash_parser import MOIS_FR, PaaResult, VolumesResult, match_paa_to_escales

# (clé, groupe, libellé) — ordre et libellés du rapport actuel
INDICATEURS = [
    ("escales", "Côte d'Ivoire", "Nb d'escales"),
    ("teu", "Côte d'Ivoire", "TEUS"),
    ("roro", "Côte d'Ivoire", "RORO"),
    ("neufs", "Côte d'Ivoire", "VEH. NEUFS"),
    ("usages", "Côte d'Ivoire", "VEH. USAGES"),
    ("t_lt15", "Côte d'Ivoire", "VEH. - 15M3"),
    ("t_15_50", "Côte d'Ivoire", "VEH. 15M3 - 50M3"),
    ("t_gt50", "Côte d'Ivoire", "VEH. + 50M3"),
    ("h_lt15", "Hinterland", "VEH. - 15M3"),
    ("h_15_50", "Hinterland", "VEH. 15M3 - 50M3"),
    ("h_gt50", "Hinterland", "VEH. + 50M3"),
    ("l_lt15", "Trafic navires Lo/Lo", "VEH. - 15M3"),
    ("l_15_50", "Trafic navires Lo/Lo", "VEH. 15M3 - 50M3"),
    ("l_gt50", "Trafic navires Lo/Lo", "VEH. + 50M3"),
]
IND_KEYS = [k for k, _, _ in INDICATEURS]
MOIS_COURT = ["Janv.", "Févr.", "Mars", "Avr.", "Mai", "Juin", "Juil.", "Août", "Sept.", "Oct.", "Nov.", "Déc."]
IND_LABEL = {k: (g, l) for k, g, l in INDICATEURS}
TRANCHE_KEYS = {"<15": "lt15", "15-50": "15_50", ">50": "gt50"}

# Sources (pastilles affichées dans l'app)
SRC_VOLUMES = "Classeur volumes"
SRC_PAA = "Extrait PAA"
SRC_SAISIE = "Saisie manuelle"
SRC_RAPPORT = "Rapport existant"
SRC_ABSENT = "À compléter"

REGLES = {
    "escales": "Nombre de navires listés dans le classeur volumes (import + export, un navire compté une fois).",
    "teu": "20' + 2 × 40', conteneurs pleins + vides, import + export. MAFI et bolsters exclus.",
    "roro": "Colonne TOTAL VEHICULE, import + export.",
    "neufs": "Colonne DT VEHICULES NEUFS, import + export.",
    "usages": "Colonne DT VEHICULES USAGES, import + export.",
    "t": "Extrait PAA, opérateur TERRA, unité VH : rubriques de déchargement (UNL) et de chargement (LOAD), "
         "transbordement compris. Shifting et bord à bord exclus.",
    "l": "Comme les tranches Côte d'Ivoire, limité aux navires de type COMBONG (Lo/Lo) dans le PAA.",
    "h": "Pas de tranche dans les sources : à saisir. Le total doit égaler la colonne DT VEH TRANSIT du classeur.",
}


def regle(ind: str) -> str:
    return REGLES.get(ind) or REGLES[ind[0]]


@dataclass
class Valeur:
    valeur: float | None
    source: str
    detail: pd.DataFrame | None = None   # contribution par navire (preuve)
    note: str = ""


# ---------------------------------------------------------------------------
# Calcul du mois
# ---------------------------------------------------------------------------
def teu_row(r) -> float:
    return r["c20_plein"] + 2 * r["c40_plein"] + r["c20_vide"] + 2 * r["c40_vide"]


def detail_par_navire(vol: VolumesResult, paa: PaaResult | None) -> pd.DataFrame:
    """Une ligne par navire : la table de preuve et le futur socle Power BI."""
    e = vol.escales.copy()
    e["teu"] = e.apply(teu_row, axis=1)
    agg = e.groupby("navire", sort=False).agg(
        armateur=("armateur", "first"), debut=("debut", "first"), fin=("fin", "first"),
        teu=("teu", "sum"), roro=("veh_total", "sum"), neufs=("veh_neufs", "sum"),
        usages=("veh_usages", "sum"), transit=("veh_transit", "sum"),
        sup50_classeur=("veh_sup50", "sum"),
        lignes_excel=("ligne_excel", lambda s: ", ".join(str(int(x)) for x in s)),
    ).reset_index()
    agg["duree_escale_h"] = (agg["fin"] - agg["debut"]).dt.total_seconds().div(3600).round(1)
    for t in ("<15", "15-50", ">50"):
        agg[f"paa_{t}"] = pd.NA
    agg["type_navire"] = ""
    agg["roro_paa"] = pd.NA
    if paa is not None:
        m = match_paa_to_escales(vol.escales, paa.lignes)
        lg = paa.lignes.assign(navire=paa.lignes["escale_paa"].map(m))
        piv = lg.dropna(subset=["navire"]).pivot_table(
            index="navire", columns="tranche", values="quantite", aggfunc="sum", fill_value=0)
        typ = lg.dropna(subset=["navire"]).groupby("navire")["type_navire"].first()
        for t in ("<15", "15-50", ">50"):
            agg[f"paa_{t}"] = agg["navire"].map(piv[t] if t in piv else {}).fillna(0)
        agg["type_navire"] = agg["navire"].map(typ).fillna("")
        agg["roro_paa"] = agg[["paa_<15", "paa_15-50", "paa_>50"]].sum(axis=1)
        agg.loc[~agg["navire"].isin(typ.index), ["paa_<15", "paa_15-50", "paa_>50", "roro_paa"]] = pd.NA
    agg["ecart_roro"] = agg["roro"] - pd.to_numeric(agg["roro_paa"], errors="coerce")
    return agg


STORE_TO_DET = {"paa_lt15": "paa_<15", "paa_15_50": "paa_15-50", "paa_gt50": "paa_>50"}


def detail_from_store(esc: pd.DataFrame) -> pd.DataFrame:
    """Détail par navire relu depuis la base -> colonnes de detail_par_navire."""
    d = esc.rename(columns=STORE_TO_DET).copy()
    for c in ["teu", "roro", "neufs", "usages", "transit", "sup50_classeur", "roro_paa",
              "paa_<15", "paa_15-50", "paa_>50", "duree_escale_h"]:
        if c in d:
            d[c] = pd.to_numeric(d[c], errors="coerce")
    d["ecart_roro"] = d["roro"] - d["roro_paa"]
    d["lignes_excel"] = ""
    return d


def compute_month(vol: VolumesResult, paa: PaaResult | None) -> tuple[dict, pd.DataFrame]:
    det = detail_par_navire(vol, paa)
    base = det[["navire", "lignes_excel"]]
    vals: dict[str, Valeur] = {}

    def from_col(col, ind):
        d = base.assign(valeur=det[col])
        vals[ind] = Valeur(float(det[col].sum()), SRC_VOLUMES, d[d["valeur"] != 0])

    vals["escales"] = Valeur(float(len(det)), SRC_VOLUMES, base.assign(valeur=1))
    from_col("teu", "teu")
    from_col("roro", "roro")
    from_col("neufs", "neufs")
    from_col("usages", "usages")

    has_paa = paa is not None and det["roro_paa"].notna().any()
    for t, suf in TRANCHE_KEYS.items():
        if has_paa:
            col = f"paa_{t}"
            d = det[["navire", "type_navire"]].assign(valeur=pd.to_numeric(det[col], errors="coerce").fillna(0))
            vals[f"t_{suf}"] = Valeur(float(d["valeur"].sum()), SRC_PAA, d[d["valeur"] != 0])
            dl = d[d["type_navire"] == "Lo/Lo"]
            vals[f"l_{suf}"] = Valeur(float(dl["valeur"].sum()), SRC_PAA, dl[dl["valeur"] != 0])
        elif t == ">50":
            d = base.assign(valeur=det["sup50_classeur"])
            vals["t_gt50"] = Valeur(float(d["valeur"].sum()), SRC_VOLUMES, d[d["valeur"] != 0],
                                    "Colonne DT SUP 50 M3 (saisie par les agents) faute d'extrait PAA.")
            vals["l_gt50"] = Valeur(None, SRC_ABSENT, note="Chargez l'extrait PAA du mois.")
        else:
            vals[f"t_{suf}"] = Valeur(None, SRC_ABSENT, note="Chargez l'extrait PAA du mois.")
            vals[f"l_{suf}"] = Valeur(None, SRC_ABSENT, note="Chargez l'extrait PAA du mois.")
        vals[f"h_{suf}"] = Valeur(None, SRC_ABSENT,
                                  note=f"À saisir. Total Hinterland du classeur : {det['transit'].sum():.0f}.")
    return vals, det


def paa_hors_classeur(vol: VolumesResult, paa: PaaResult | None) -> list[str]:
    """Escales PAA TERRA sans navire correspondant dans le classeur volumes."""
    if paa is None:
        return []
    ok = set(match_paa_to_escales(vol.escales, paa.lignes))
    out = []
    for esc, g in paa.lignes.groupby("escale_paa"):
        if esc not in ok:
            out.append(f"{g['navire_paa'].iloc[0]} (escale {esc}, {g['type_navire'].iloc[0]}) : "
                       f"{g['quantite'].sum():.0f} véhicule(s) au PAA, absent du classeur.")
    return out


def compute_paa_only(paa: PaaResult) -> dict:
    """Mois sans classeur de volumes (ex. 2025) : seules les tranches de volume et le
    trafic Lo/Lo se déduisent de l'extrait PAA. Les autres indicateurs ne sont pas touchés."""
    vals: dict[str, Valeur] = {}
    lg = paa.lignes
    for t, suf in TRANCHE_KEYS.items():
        d = lg[lg["tranche"] == t]
        vals[f"t_{suf}"] = Valeur(float(d["quantite"].sum()), SRC_PAA, d[["navire_paa", "type_navire", "quantite"]])
        dl = d[d["type_navire"] == "Lo/Lo"]
        vals[f"l_{suf}"] = Valeur(float(dl["quantite"].sum()), SRC_PAA, dl[["navire_paa", "type_navire", "quantite"]])
    return vals


def controles(det: pd.DataFrame, retenu: dict, alertes=(), periode_paa: tuple | None = None,
              periode: tuple | None = None, alertes_paa=()) -> pd.DataFrame:
    """det = détail par navire (colonnes de detail_par_navire) ;
    retenu = {indicateur: valeur retenue (corrections comprises)} ;
    alertes = messages du lecteur de fichier ; periode(_paa) = (annee, mois)."""
    rows = []

    def add(nom, rapport, ref, explication, tol=0):
        if rapport is None or ref is None or pd.isna(rapport) or pd.isna(ref):
            rows.append([nom, rapport, ref, None, "À compléter", explication])
            return
        ecart = rapport - ref
        statut = "OK" if abs(ecart) <= tol else "À vérifier"
        rows.append([nom, rapport, ref, ecart, statut, explication])

    g = lambda k: retenu.get(k)
    add("RORO = neufs + usagés", g("roro"), (g("neufs") or 0) + (g("usages") or 0),
        "Si l'écart n'est pas nul, une ligne navire a un total différent de neufs + usagés.")
    tr = [g(k) for k in ("t_lt15", "t_15_50", "t_gt50")]
    add("RORO = somme des 3 tranches", g("roro"), sum(tr) if None not in tr else None,
        "Compare le total du classeur aux tranches du PAA. Voir le détail par navire pour localiser l'écart.")
    ht = [g(k) for k in ("h_lt15", "h_15_50", "h_gt50")]
    add("Hinterland : somme des tranches = transit du classeur", sum(ht) if None not in ht else None,
        float(det["transit"].sum()), "Les tranches Hinterland sont saisies : leur total doit égaler DT VEH TRANSIT.")
    if "roro_paa" in det and pd.to_numeric(det["roro_paa"], errors="coerce").notna().any():
        det = det.assign(roro_paa=pd.to_numeric(det["roro_paa"], errors="coerce"))
        det = det.assign(ecart_roro=det["roro"] - det["roro_paa"])
        d = det.dropna(subset=["roro_paa"])
        add("Tranche + 50 m³ : PAA vs DT SUP 50 M3 saisi", float(pd.to_numeric(d["paa_>50"]).sum()),
            float(d["sup50_classeur"].sum()),
            "La colonne DT SUP 50 M3 est saisie à la main ; le PAA applique les rubriques de facturation.")
        for _, r in d[d["ecart_roro"].fillna(0) != 0].iterrows():
            rows.append([f"Navire {r['navire']} : RORO classeur vs PAA", r["roro"], r["roro_paa"],
                         r["ecart_roro"], "À vérifier", "Écart de comptage sur ce navire entre les deux sources."])
        sans = det[det["roro_paa"].isna() & (det["roro"] > 0)]["navire"].tolist()
        if sans:
            rows.append(["Navires absents du PAA", None, None, len(sans), "À vérifier", ", ".join(sans)])
        if periode_paa and periode and periode_paa[1] and tuple(periode_paa) != tuple(periode):
            rows.append(["Mois du PAA = mois du classeur", None, None, None, "À vérifier",
                         f"PAA : {MOIS_FR[periode_paa[1] - 1]} {periode_paa[0]} ; "
                         f"classeur : {MOIS_FR[periode[1] - 1]} {periode[0]}."])
    for a in alertes_paa:
        rows.append(["PAA : point à vérifier", None, None, None, "À vérifier", a])
    for a in alertes:
        rows.append(["Classeur : lignes navires = ligne TOTAL", None, None, None, "À vérifier", a])
    return pd.DataFrame(rows, columns=["Contrôle", "Valeur rapport", "Valeur de contrôle", "Écart", "Statut", "Explication"])


# ---------------------------------------------------------------------------
# Référentiel : lecture du rapport existant (amorçage)
# ---------------------------------------------------------------------------
_ROW_BY_IND = dict(zip(IND_KEYS, [48, 49, 50, 52, 53, 54, 55, 56, 58, 59, 60, 62, 63, 64]))


def _safe_eval(v):
    """Évalue une formule purement arithmétique (=3247*0.43, =(32826/12)*9)."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).lstrip("=").replace(",", ".").strip()
    if not re.fullmatch(r"[\d\.\s\+\-\*/\(\)]+", s):
        return None
    try:
        return float(eval(s, {"__builtins__": {}}, {}))  # noqa: S307 — chiffres et opérateurs uniquement
    except Exception:
        return None


def _annual_from_prorata(v):
    """Q48 '=(150/12)*9' -> 150 (total annuel 2025)."""
    if isinstance(v, (int, float)) and v == 0:
        return 0.0
    m = re.match(r"=\(?\s*([\d\.]+)\s*/\s*12\s*\)?\s*\*\s*\d+", str(v or ""))
    if m:
        return float(m.group(1))
    return None


def parse_rapport_existant(data: bytes) -> dict:
    """Lit le bloc « REPORTING RORO & TEUS » du rapport actuel (lignes 47-64).
    Retourne {"annee", "realise": {(mois, ind): v}, "annuel_2025": {ind: v},
    "mois_ref_2025": int, "meme_mois_2025": {ind: v}, "budget": {ind: v}}."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = None
    for w in wb.worksheets:
        for r in range(40, 50):
            if str(w.cell(r, 8).value or "").strip().upper().startswith("REPORTING RORO"):
                ws = w
                break
        if ws:
            break
    if ws is None:
        raise ValueError("Bloc « REPORTING RORO & TEUS » introuvable (attendu vers la ligne 45, colonne H).")
    head = 47
    mois_cols = {}
    for c in range(4, 16):
        v = ws.cell(head, c).value
        if hasattr(v, "month"):
            mois_cols[v.month] = c
            annee = v.year
    out = {"annee": annee, "realise": {}, "annuel_2025": {}, "meme_mois_2025": {}, "budget": {}}
    lib_v = str(ws.cell(head, 22).value or "")
    mref = next((i + 1 for i, m in enumerate(MOIS_FR) if m.upper().replace("Û", "U") in lib_v.upper().replace("Û", "U")), None)
    out["mois_ref_2025"] = mref
    for ind, r in _ROW_BY_IND.items():
        for m, c in mois_cols.items():
            v = _safe_eval(ws.cell(r, c).value)
            if v is not None:
                out["realise"][(m, ind)] = v
        q = ws.cell(r, 17).value
        a = _annual_from_prorata(q)
        if a is not None:
            out["annuel_2025"][ind] = a
        b = _safe_eval(ws.cell(r, 19).value)
        if b is not None:
            out["budget"][ind] = b
        vv = _safe_eval(ws.cell(r, 22).value)
        if vv is not None:
            out["meme_mois_2025"][ind] = vv
    return out


# ---------------------------------------------------------------------------
# Tableau mensuel (affichage + export)
# ---------------------------------------------------------------------------
def _pct(a, b):
    if a is None or b in (None, 0) or pd.isna(a) or pd.isna(b):
        return None
    return (a - b) / b


def monthly_table(realise26: dict, realise25: dict, annuel25: dict, budget: dict, annee: int, n: int) -> pd.DataFrame:
    """realise26/realise25 = {(mois, ind): v} ; annuel25/budget = {ind: v} ; n = mois de référence."""
    rows = []
    for ind, grp, lib in INDICATEURS:
        r = {"Groupe": grp, "Indicateur": lib, "_ind": ind}
        for m in range(1, 13):
            r[MOIS_COURT[m - 1]] = realise26.get((m, ind))
        # cumul sur les seuls mois disponibles en N ; N-1 et budget comparés sur ces mêmes mois
        dispo = [m for m in range(1, n + 1) if realise26.get((m, ind)) is not None]
        k = len(dispo)
        tot26 = sum(realise26[(m, ind)] for m in dispo) if k else None
        v25 = [realise25.get((m, ind)) for m in dispo]
        if k and all(v is not None for v in v25):
            tot25, base25 = sum(v25), "réel"
        elif k and annuel25.get(ind) is not None:
            tot25, base25 = annuel25[ind] / 12 * k, "proratisé"
        else:
            tot25, base25 = None, ""
        bud = budget.get(ind)
        cur = realise26.get((n, ind))
        same25 = realise25.get((n, ind))
        r.update({
            f"Total {annee} ({n} mois)": tot26,
            f"Total {annee - 1} ({n} mois)": tot25,
            "_base25": base25,
            "_mois_cumules": k,
            "Budget / mois": bud,
            f"Cumul budget ({n} mois)": bud * k if bud is not None and k else None,
            f"{MOIS_FR[n - 1].capitalize()} {annee - 1}": same25,
            "% mois R/B": _pct(cur, bud),
            "% cumul R/B": _pct(tot26, bud * k if bud is not None and k else None),
            f"% mois {annee}/{annee - 1}": _pct(cur, same25),
            f"% cumul {annee}/{annee - 1}": _pct(tot26, tot25),
        })
        rows.append(r)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Export Excel : mise en page du rapport + formules vivantes + preuves
# ---------------------------------------------------------------------------
def build_export(realise26, realise25, annuel25, budget, annee, n, sources: dict,
                 detail: pd.DataFrame | None, ctrl: pd.DataFrame | None, corrections: pd.DataFrame | None) -> bytes:
    import xlsxwriter
    from xlsxwriter.utility import xl_rowcol_to_cell as cell

    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    F = "Segoe UI"
    title = wb.add_format({"font_name": F, "bold": True, "font_size": 14, "font_color": "#1F3864"})
    hdr = wb.add_format({"font_name": F, "bold": True, "bg_color": "#1F3864", "font_color": "white",
                         "border": 1, "align": "center", "valign": "vcenter", "text_wrap": True})
    grp = wb.add_format({"font_name": F, "bold": True, "bg_color": "#D9E1F2", "border": 1})
    lab = wb.add_format({"font_name": F, "border": 1})
    num = wb.add_format({"font_name": F, "border": 1, "num_format": "#,##0"})
    num_b = wb.add_format({"font_name": F, "border": 1, "num_format": "#,##0", "bold": True, "bg_color": "#F2F2F2"})
    num_corr = wb.add_format({"font_name": F, "border": 1, "num_format": "#,##0", "bg_color": "#FFF2CC"})
    pct = wb.add_format({"font_name": F, "border": 1, "num_format": "0.0%;[Red]-0.0%;0.0%", "align": "right"})
    note = wb.add_format({"font_name": F, "italic": True, "font_color": "#595959"})
    wrap = wb.add_format({"font_name": F, "text_wrap": True, "valign": "top", "border": 1})

    ws = wb.add_worksheet("REPORTING RORO & TEUS")
    ws.write(0, 0, f"REPORTING RORO & TEUS — {annee} (arrêté à fin {MOIS_FR[n - 1]})", title)
    ws.write(1, 0, "Cellules jaunes : valeur corrigée à la main (voir onglet « Corrections »). "
                   "« — » : division impossible (budget ou référence absent).", note)
    H = 3
    heads = ["", "Indicateur"] + MOIS_COURT + [
        f"TOTAL {annee} ({n} mois)", f"TOTAL {annee - 1} ({n} mois)", "Budget / mois",
        f"Cumul budget ({n} mois)", f"{MOIS_FR[n - 1].capitalize()} {annee - 1}",
        "% R/B mois", "% cumul R/B", f"% R{annee % 100}/R{(annee - 1) % 100} mois",
        f"% cumul R{annee % 100}/R{(annee - 1) % 100}", "Source du mois"]
    for c, h in enumerate(heads):
        ws.write(H, c, h, hdr)
    ws.set_row(H, 36)
    ws.set_column(0, 0, 20)
    ws.set_column(1, 1, 18)
    ws.set_column(2, 13, 8)
    ws.set_column(14, 23, 12)
    ws.set_column(24, 24, 20)
    r = H + 1
    last_grp = None
    corr_keys = set()
    if corrections is not None and not corrections.empty:
        corr_keys = {(int(a), int(m), i) for a, m, i in corrections[["annee", "mois", "indicateur"]].itertuples(index=False)}
    for ind, g, l in INDICATEURS:
        if g != last_grp:
            ws.merge_range(r, 0, r, len(heads) - 1, g.upper(), grp)
            r += 1
            last_grp = g
        ws.write(r, 0, "", lab)
        ws.write(r, 1, l, lab)
        for m in range(1, 13):
            v = realise26.get((m, ind))
            fmt = num_corr if (annee, m, ind) in corr_keys else num
            if v is None:
                ws.write_blank(r, 1 + m, None, fmt)
            else:
                ws.write_number(r, 1 + m, v, fmt)
        c_tot, c_25, c_bud, c_cbud, c_same = 14, 15, 16, 17, 18
        c_m = 1 + n
        ws.write_formula(r, c_tot, f"=SUM({cell(r, 2)}:{cell(r, 1 + n)})", num_b)
        dispo = [m for m in range(1, n + 1) if realise26.get((m, ind)) is not None]
        k = len(dispo)  # cumul sur les mois disponibles en N ; N-1 et budget sur les mêmes mois
        v25 = [realise25.get((m, ind)) for m in dispo]
        if k and all(v is not None for v in v25):
            ws.write_number(r, c_25, sum(v25), num)
        elif k and annuel25.get(ind) is not None:
            ws.write_formula(r, c_25, f"=({annuel25[ind]:g}/12)*{k}", num)
        else:
            ws.write_blank(r, c_25, None, num)
        if budget.get(ind) is not None:
            ws.write_number(r, c_bud, budget[ind], num)
            ws.write_formula(r, c_cbud, f"={cell(r, c_bud)}*{k}", num)
        else:
            ws.write_blank(r, c_bud, None, num)
            ws.write_blank(r, c_cbud, None, num)
        same = realise25.get((n, ind))
        if same is not None:
            ws.write_number(r, c_same, same, num)
        else:
            ws.write_blank(r, c_same, None, num)

        def pf(a, b):
            return f'=IF(OR({b}="",{b}=0,{a}=""),"—",({a}-{b})/{b})'
        ws.write_formula(r, 19, pf(cell(r, c_m), cell(r, c_bud)), pct)
        ws.write_formula(r, 20, pf(cell(r, c_tot), cell(r, c_cbud)), pct)
        ws.write_formula(r, 21, pf(cell(r, c_m), cell(r, c_same)), pct)
        ws.write_formula(r, 22, pf(cell(r, c_tot), cell(r, c_25)), pct)
        ws.write(r, 24, sources.get(ind, ""), lab)
        r += 1
    ws.freeze_panes(H + 1, 2)

    if detail is not None and not detail.empty:
        wd = wb.add_worksheet("Détail par navire")
        cols = ["navire", "type_navire", "armateur", "debut", "fin", "duree_escale_h", "teu", "roro",
                "neufs", "usages", "transit", "paa_<15", "paa_15-50", "paa_>50", "roro_paa", "ecart_roro",
                "sup50_classeur", "lignes_excel"]
        labels = ["Navire", "Type", "Armateur", "Début opérations", "Fin opérations", "Durée (h)", "TEU",
                  "RORO", "Neufs", "Usagés", "Hinterland", "PAA <15", "PAA 15-50", "PAA >50", "RORO PAA",
                  "Écart classeur - PAA", "DT SUP 50 M3 (classeur)", "Lignes du classeur"]
        dfmt = wb.add_format({"font_name": F, "border": 1, "num_format": "dd/mm/yyyy hh:mm"})
        for c, h in enumerate(labels):
            wd.write(0, c, h, hdr)
        for i, row in enumerate(detail[cols].itertuples(index=False), start=1):
            for c, v in enumerate(row):
                if v is None or (not isinstance(v, str) and pd.isna(v)):
                    wd.write_blank(i, c, None, lab)
                elif c in (3, 4):
                    wd.write_datetime(i, c, pd.Timestamp(v).to_pydatetime(), dfmt)
                elif isinstance(v, (int, float)):
                    wd.write_number(i, c, float(v), num)
                else:
                    wd.write(i, c, str(v), lab)
        wd.set_column(0, 0, 28)
        wd.set_column(1, 2, 16)
        wd.set_column(3, 4, 17)
        wd.set_column(5, 17, 11)
        wd.freeze_panes(1, 1)
        wd.autofilter(0, 0, len(detail), len(cols) - 1)

    if ctrl is not None and not ctrl.empty:
        wc = wb.add_worksheet("Contrôles")
        for c, h in enumerate(ctrl.columns):
            wc.write(0, c, h, hdr)
        for i, row in enumerate(ctrl.itertuples(index=False), start=1):
            for c, v in enumerate(row):
                if v is None or (not isinstance(v, str) and pd.isna(v)):
                    wc.write_blank(i, c, None, lab)
                elif isinstance(v, (int, float)):
                    wc.write_number(i, c, float(v), num)
                else:
                    wc.write(i, c, str(v), wrap)
        wc.set_column(0, 0, 44)
        wc.set_column(1, 4, 14)
        wc.set_column(5, 5, 70)

    if corrections is not None and not corrections.empty:
        wk = wb.add_worksheet("Corrections")
        cc = corrections.copy()
        cc["indicateur"] = cc["indicateur"].map(lambda k: " · ".join(IND_LABEL.get(k, ("", k))))
        for c, h in enumerate(cc.columns):
            wk.write(0, c, h, hdr)
        for i, row in enumerate(cc.itertuples(index=False), start=1):
            for c, v in enumerate(row):
                if v is None or (not isinstance(v, str) and pd.isna(v)):
                    wk.write_blank(i, c, None, lab)
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    wk.write_number(i, c, float(v), lab)
                else:
                    wk.write(i, c, str(v), lab)
        wk.set_column(0, len(cc.columns), 18)

    wr = wb.add_worksheet("Sources & règles")
    wr.write(0, 0, "Sources et règles de calcul", title)
    wr.write(2, 0, "Indicateur", hdr)
    wr.write(2, 1, "Règle", hdr)
    wr.write(2, 2, f"Source {MOIS_FR[n - 1]} {annee}", hdr)
    for i, (ind, g, l) in enumerate(INDICATEURS, start=3):
        wr.write(i, 0, f"{g} · {l}", lab)
        wr.write(i, 1, regle(ind), wrap)
        wr.write(i, 2, sources.get(ind, ""), lab)
    wr.set_column(0, 0, 34)
    wr.set_column(1, 1, 90)
    wr.set_column(2, 2, 34)
    wb.close()
    return buf.getvalue()
